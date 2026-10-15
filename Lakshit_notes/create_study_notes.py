from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'Data_Analytics_Combined_Study_Notes.docx'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def borders(table):
    b = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),'D9D9D9'); b.append(e)
    table._tbl.tblPr.append(b)

def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'1F4E78')
        for r in c.paragraphs[0].runs: r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(255,255,255)
    tr=t.rows[0]._tr.get_or_add_trPr(); tr.append(OxmlElement('w:tblHeader'))
    for ri,row in enumerate(rows):
        cells=t.add_row().cells; tr=t.rows[-1]._tr.get_or_add_trPr(); tr.append(OxmlElement('w:cantSplit'))
        for i,val in enumerate(row):
            cells[i].text=str(val)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs: r.font.size=Pt(9)
            if ri%2: shade(cells[i],'F3F6F9')
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def bullet(doc, text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(text)

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
for st in ['Normal','List Bullet']:
    doc.styles[st].font.name='Aptos'; doc.styles[st].font.size=Pt(10.5); doc.styles[st]._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos')
for n,sz in [('Title',22),('Heading 1',15),('Heading 2',12)]:
    doc.styles[n].font.name='Aptos'; doc.styles[n].font.size=Pt(sz); doc.styles[n].font.bold=True; doc.styles[n].font.color.rgb=RGBColor(0,0,0)

p=doc.add_paragraph(style='Title'); p.add_run('Data Analytics Combined Study Notes')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(12); r=p.add_run('Integrated notes from the syllabus and Day1 to Day10 folders'); r.italic=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor(79,79,79)

doc.add_heading('How to use these notes',1)
doc.add_paragraph('These notes combine the practical material found in the Day1–Day10 folders with the complete syllabus in 03_Data Analytics.pdf. Folder topics are practical coverage; syllabus-only topics are included as study areas to complete. The syllabus is arranged in Sessions 1–25, while the folders are arranged by day, so the sequence follows concepts rather than forcing an exact day-to-session match.')

doc.add_heading('1 Foundations and the analytics life cycle',1)
doc.add_paragraph('Business analytics uses data to describe what happened, explain why it happened, predict what may happen, and support decisions. The analytics life cycle is:')
for x in ['Discovery: define the question, stakeholders, constraints, and success measure.', 'Data preparation: acquire data, inspect its schema, clean errors, handle missing values, and document assumptions.', 'Model planning: choose variables, methods, validation strategy, and evaluation metrics.', 'Model building: implement the analysis or model and compare alternatives.', 'Quality assurance and documentation: check data quality, reproducibility, and interpretation.', 'Management approval, installation, acceptance, and operation: deliver the result into a usable workflow and monitor it.']:
    bullet(doc,x)
doc.add_paragraph('Folder connection: Day1 supplies advanced Python foundations. Day2 supplies NumPy and Pandas foundations. The business life-cycle activities themselves are mainly syllabus topics and are not fully represented in the folders.')

doc.add_heading('2 Python foundations from Day1',1)
add_table(doc,['Topic','Key notes'],[
['Closures and nested functions','Nested functions can retain variables from an enclosing scope.'],
['Decorators and __call__','Decorators wrap functions or classes; __call__ makes an object callable.'],
['Polymorphism, inheritance and mixins','Common interfaces allow interchangeable objects; mixins add focused behavior.'],
['Comprehensions and generators','Comprehensions build collections; generators yield values lazily and save memory.'],
['Coroutines and context managers','Coroutines receive values over time; context managers guarantee setup and cleanup.'],
['Threads and scope','Threads support concurrent I/O work; understand local, global, nonlocal, and enclosing scope.'],
['Partial functions and currying','Pre-fill selected arguments to create simpler functions.'],
])

doc.add_heading('3 NumPy and Pandas from Day2',1)
doc.add_paragraph('NumPy provides efficient n-dimensional arrays. Pandas provides labelled Series and DataFrames.')
for x in ['ndarray and tensor initialization; shape, dtype, indexing, slicing, and selection.', 'Vectorized operations and broadcasting apply computations without explicit Python loops.', 'Linear algebra includes matrix multiplication, decompositions, and solving systems.', 'Pandas Series stores labelled values; DataFrame stores labelled rows and columns.', 'Inspect head, info, describe, shape, dtypes, and value_counts before analysis.']:
    bullet(doc,x)

doc.add_heading('4 Descriptive statistics from Days3, 4, 6 and 7',1)
add_table(doc,['Measure group','What to know'],[
['Central tendency','Mean, median, and mode describe the centre. Geometric and harmonic means apply to positive data.'],
['Dispersion','Range, variance, standard deviation, IQR, mean absolute deviation, median absolute deviation, and coefficient of variation describe spread.'],
['Position','Quartiles, percentiles, percentile rank, and z-scores describe relative standing.'],
['Shape','Skewness measures asymmetry; kurtosis measures tail weight. Higher moments are sensitive to outliers.'],
['Robustness','Median and IQR are generally more resistant to extreme values than mean and standard deviation.'],
])
doc.add_paragraph('Syllabus connection: Session 5 covers descriptive statistics and Sessions 11–12 name skewness. The folders also include detailed percentile methods and distribution-shape calculations.')

doc.add_heading('5 Data population, probability and distributions',1)
for x in ['Population is the full set of units; a sample is a subset used to learn about the population.', 'A sample space contains all possible outcomes; an event is a subset of outcomes.', 'Joint probability concerns two events together. Conditional probability updates probability given information. Marginal probability describes one variable after ignoring or summing over another.', 'Bayes theorem updates a prior belief using evidence: P(A|B) = P(B|A)P(A) / P(B).', 'Random variables may be discrete or continuous. Normal, binomial, and Poisson distributions model common situations.', 'The Central Limit Theorem explains why sample means often approach a normal distribution as sample size grows.', 'Sampling and estimation use sample statistics to estimate population parameters and quantify uncertainty.']:
    bullet(doc,x)
doc.add_paragraph('Days5 and 6 contain data-generation and distribution exercises. The syllabus additionally requires sample spaces, Bayes theorem, sampling, estimation, and the Central Limit Theorem.')

doc.add_heading('6 Correlation, covariance and outliers',1)
for x in ['Covariance indicates the direction of joint variation but depends on measurement units.', 'Pearson correlation standardizes covariance to a value between -1 and +1 and measures linear association.', 'Spearman correlation uses ranks and measures monotonic association. Kendall correlation compares concordant and discordant pairs.', 'Correlation does not establish causation. Check plots, confounders, sample size, and data quality.', 'Outliers may be errors, rare events, or genuine extremes. Inspect, document, and choose a defensible treatment such as correction, transformation, robust statistics, clipping, or a separate model.']:
    bullet(doc,x)

doc.add_heading('7 Data preparation and exploratory data analysis',1)
for x in ['Indexing and selection: use label-based and position-based selection deliberately; verify alignment.', 'Missing and malformed data: identify missingness, invalid types, duplicates, inconsistent labels, and impossible values.', 'Grouping and reshaping: aggregate by categories; use pivot, melt, stack, and unstack.', 'Large data I/O: use efficient formats, chunks, appropriate dtypes, and memory-aware operations.', 'EDA workflow: inspect schema, distributions, counts, relationships, missingness, and unusual observations.', 'Visualization: use Matplotlib and Seaborn for distributions, comparisons, relationships, and diagnostic plots.', 'Profiling tools automate summaries of missingness, duplicates, cardinality, distributions, and correlations; verify important findings manually.']:
    bullet(doc,x)
doc.add_paragraph('Day9 focuses on Pandas wrangling. Day10 focuses on EDA and profiling. These support the syllabus topics of data preparation, intelligent data analysis, exploring data, and visualization.')

doc.add_heading('8 Hypothesis testing and statistical inference',1)
doc.add_paragraph('These syllabus topics are not clearly represented in the Day1–Day10 folders.')
for x in ['State a null hypothesis and alternative hypothesis before inspecting results.', 'A p-value measures how surprising the observed result would be if the null hypothesis were true.', 'Use a Z-test when assumptions about the sampling distribution and variance are appropriate.', 'A Chi-square test is commonly used for categorical association or goodness of fit.', 'Report effect size, confidence intervals, assumptions, and practical meaning with statistical significance.']:
    bullet(doc,x)

doc.add_heading('9 Predictive and prescriptive modelling',1)
for x in ['Predictive modelling estimates an outcome from informative attributes. Classification predicts categories; regression predicts numerical values.', 'Segmentation divides observations into meaningful groups. Supervised segmentation uses a target to guide the split.', 'Decision trees express learned models as a sequence of rules; inspect depth, interpretability, and generalization.', 'Probability estimation communicates uncertainty rather than only a hard class label.', 'Prescriptive modelling recommends an action under constraints by connecting predictions with objectives, costs, risks, and feasible decisions.']:
    bullet(doc,x)

doc.add_heading('10 Regression and forecasting',1)
for x in ['Regression models a numerical response. Examine coefficients, residuals, fit, assumptions, and multicollinearity.', 'Forecasting uses ordered time observations. Separate time correctly, avoid leakage from the future, and evaluate with suitable error measures.', 'Compare every model with a simple baseline. Complexity is useful only when it improves reliable out-of-sample performance.']:
    bullet(doc,x)

doc.add_heading('11 Simulation, risk and optimization',1)
for x in ['Simulation generates plausible outcomes under specified assumptions and explores uncertainty.', 'Risk analysis summarizes adverse outcomes, probabilities, sensitivity, and consequences.', 'Linear optimization uses linear objectives and constraints; nonlinear optimization allows curved relationships and may require different solvers.', 'State assumptions, bounds, units, and the decision objective.']:
    bullet(doc,x)

doc.add_heading('12 Validation and decision analytics',1)
for x in ['Overfitting occurs when a model learns noise. Use simpler models, regularization, more data, or better features.', 'Generalization is performance on unseen data. Holdout evaluation uses a reserved set; cross-validation rotates validation folds.', 'Evaluate classifiers with a confusion matrix, accuracy, precision, recall, F1-score, ROC or precision-recall curves, and class-aware baselines.', 'Decision analytics connects evidence, model outputs, costs, benefits, and implications for investment in data.']:
    bullet(doc,x)

doc.add_heading('13 Bayesian reasoning and evidence',1)
for x in ['Separate prior belief, evidence, likelihood, posterior belief, and the assumptions connecting them.', 'Combine explicit evidence with Bayes rule while checking independence assumptions and base rates.', 'Communicate probabilistic conclusions with uncertainty and sensitivity to assumptions.']:
    bullet(doc,x)

doc.add_heading('14 Factor, directional and functional data analysis',1)
for x in ['Factor analysis explains correlations through a smaller number of latent factors.', 'Directional data has circular or spherical structure; ordinary linear summaries may be misleading.', 'Functional data analysis treats curves or signals as observations and studies variation over a continuous domain.']:
    bullet(doc,x)

doc.add_heading('15 Power BI and communication',1)
for x in ['Power BI connects data sources, Power Query transformations, a model, visuals, reports, dashboards, and sharing permissions.', 'Power Query imports, cleans, reshapes, and combines data before modelling.', 'Good dashboards focus on a decision, use clear comparisons, provide context, and make definitions and refresh timing visible.', 'A visual should help a reader compare, discover a relationship, understand change, or act.']:
    bullet(doc,x)

doc.add_heading('16 Practical study checklist',1)
add_table(doc,['Stage','Questions to answer'],[
['Prepare','What is the question? What are the units, variables, types, missing values, and quality risks?'],
['Explore','What do distributions, groups, plots, and relationships show? What is unusual?'],
['Measure','Which descriptive or inferential measure is appropriate, and what assumptions does it require?'],
['Model','What is the target, baseline, validation design, metric, and acceptable error?'],
['Decide','What action follows? What are the costs, risks, uncertainty, and constraints?'],
['Communicate','Can another person reproduce the work and understand the conclusion?'],
])
doc.add_heading('Topics to complete after Day10',1)
doc.add_paragraph('The major remaining syllabus areas are hypothesis testing, predictive and prescriptive modelling, regression, forecasting, simulation, optimization, validation and classifier evaluation, Bayesian reasoning, factor/directional/functional analysis, and Power BI. The folders provide a strong base in programming, arrays, tabular data, descriptive statistics, probability distributions, correlation, wrangling, and EDA.')
f=doc.add_paragraph(); f.paragraph_format.space_before=Pt(8); r=f.add_run('Prepared from 03_Data Analytics.pdf and the Day1–Day10 folder contents. Review date: 15 October 2026.'); r.italic=True; r.font.color.rgb=RGBColor(79,79,79)
doc.save(OUT); print(OUT)
